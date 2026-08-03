#!/usr/bin/env python3
"""基本設計書（外部設計書）生成"""

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference, BarChart
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment

from generate_design_docs import (
    OUT_DIR, thin, fill_navy, fill_teal, fill_section, fill_light, fill_alt,
    fill_white, fill_warn, fill_ok, fill_accent, font_h2, font_body, font_small,
    font_label, align_l, align_t, align_c, NAVY, TEAL, ACCENT, WHITE,
    set_col_widths, merge_write, header_bar, section_bar, table_header,
    table_row, write_cover, write_toc,
)


def build_basic_design():
    wb = Workbook()
    ws = wb.active
    ws.title = "00_表紙"
    write_cover(ws, "基本設計書（外部設計書）", "DOC-JLMS-BD-001", "1.0", "正式版（Approved）")

    ws = wb.create_sheet("01_目次")
    toc = [
        ("1", "はじめに・設計方針", "02_設計方針", "—", "アーキ", "確定", "◎", ""),
        ("2", "システム構成", "03_システム構成", "—", "アーキ", "確定", "◎", ""),
        ("3", "機能構成・モジュール", "04_機能構成", "—", "設計", "確定", "◎", ""),
        ("4", "画面遷移設計", "05_画面遷移", "—", "設計", "確定", "◎", ""),
        ("5", "画面一覧・レイアウト方針", "06_画面設計", "—", "設計", "確定", "◎", ""),
        ("6", "帳票設計", "07_帳票設計", "—", "設計", "確定", "◎", ""),
        ("7", "論理データモデル", "08_データモデル", "—", "DBA", "確定", "◎", ""),
        ("8", "テーブル一覧", "09_テーブル一覧", "—", "DBA", "確定", "◎", ""),
        ("9", "外部IF設計", "10_IF設計", "—", "設計", "確定", "◎", ""),
        ("10", "バッチ設計", "11_バッチ設計", "—", "設計", "確定", "◎", ""),
        ("11", "セキュリティ設計", "12_セキュリティ", "—", "アーキ", "確定", "◎", ""),
        ("12", "運用・監視設計", "13_運用監視", "—", "インフラ", "確定", "◎", ""),
        ("13", "要件トレーサビリティ", "14_トレーサビリティ", "—", "QA", "確定", "◎", ""),
    ]
    write_toc(ws, toc, "基本設計書　目次")

    # ── 設計方針 ──
    ws = wb.create_sheet("02_設計方針")
    set_col_widths(ws, [12, 20, 36, 14, 14, 12, 12, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "1. 設計方針")
    section_bar(ws, 3, 8, "1.1 目的")
    merge_write(
        ws, "A4", "H5",
        "本基本設計書は、要件定義書（DOC-JLMS-RD-001 版1.0）を実現するためのシステム方式・機能配置・"
        "画面／データ／IFの外部仕様を定義し、詳細設計・実装のインプットとする。",
        fill=fill_white, font=font_body, align=align_t, border=thin,
    )
    section_bar(ws, 7, 8, "1.2 アーキテクチャ原則")
    table_header(ws, 8, ["ID", "原則", "内容", "根拠要件", "トレードオフ", "決定", "例外", "備考"])
    principles = [
        ["AP-01", "API First", "UIとAPIを分離しOpenAPIで契約", "保守性", "初期工数増", "採用", "なし", ""],
        ["AP-02", "論理マルチテナント", "共有DB＋tenant_id強制フィルタ", "FR-SYS-01", "物理分離は将来", "採用", "なし", ""],
        ["AP-03", "イベント収集", "学習操作をイベントとして永続化", "FR-LRN-01", "容量増", "採用", "なし", "分析基盤"],
        ["AP-04", "マネージド優先", "AWSマネージドサービスを優先採用", "NFR運用", "ベンダーロック", "採用", "なし", ""],
        ["AP-05", "ゼロトラスト寄り", "全API認証必須・最小権限・監査", "NFR-06〜", "開発摩擦", "採用", "ヘルスチェックのみ公開", ""],
        ["AP-06", "段階リリース", "Feature Flagで機能段階公開", "リスク", "複雑性", "採用", "基盤機能", ""],
    ]
    for i, t in enumerate(principles):
        table_row(ws, 9 + i, t, alt=i % 2 == 1, heights=22)

    section_bar(ws, 17, 8, "1.3 技術スタック（決定事項）")
    table_header(ws, 18, ["層", "技術", "版目安", "選定理由", "代替案", "決定日", "責任者", "備考"])
    stack = [
        ["フロント", "React + TypeScript", "18.x", "コンポーネント再利用・型安全", "Vue", "2026-04", "設計", "Vite"],
        ["BFF/API", "Node.js (NestJS)", "20.x LTS", "OpenAPI親和・生産性", "Java Spring", "2026-04", "アーキ", ""],
        ["ドメインDB", "Amazon Aurora PostgreSQL", "15.x", "運用性・互換", "RDS PG", "2026-04", "DBA", "Multi-AZ"],
        ["キャッシュ", "Amazon ElastiCache Redis", "7.x", "セッション・レート制限", "MemoryDB", "2026-04", "アーキ", ""],
        ["オブジェクト", "S3 + CloudFront", "—", "教材配信・署名URL", "—", "2026-04", "インフラ", ""],
        ["認証", "Amazon Cognito連携 or 直接OIDC", "—", "Entra IDフェデレーション", "Auth0", "2026-05", "アーキ", "OIDC"],
        ["バッチ", "AWS Batch / EventBridge", "—", "スケジュール・再試行", "ECS Scheduled", "2026-05", "設計", ""],
        ["監視", "CloudWatch + X-Ray", "—", "統合運用", "Datadog", "2026-05", "運用", "APM"],
        ["CI/CD", "GitHub Actions → ECR → ECS", "—", "既存開発フロー", "CodePipeline", "2026-04", "開発", ""],
    ]
    for i, t in enumerate(stack):
        table_row(ws, 19 + i, t, alt=i % 2 == 1, heights=20)

    # ── システム構成 ──
    ws = wb.create_sheet("03_システム構成")
    set_col_widths(ws, [14, 18, 28, 16, 14, 14, 12, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "2. システム構成")
    section_bar(ws, 3, 8, "2.1 論理構成")
    merge_write(
        ws, "A4", "H8",
        "【クライアント層】　ブラウザ（Responsive Web）\n"
        "　　↓ HTTPS / CloudFront\n"
        "【プレゼンテーション】　静的ホスト（S3+CloudFront）＋ React SPA\n"
        "　　↓ HTTPS / JWT\n"
        "【アプリケーション層】　ECS Fargate（APIサービス）× N　／　非同期ワーカー\n"
        "　　↓\n"
        "【データ層】　Aurora PostgreSQL　／　ElastiCache　／　S3（教材・帳票）\n"
        "【横断】　Entra ID（OIDC）　／　SES（メール）　／　EventBridge（スケジュール）　／　CloudWatch",
        fill=fill_light, font=font_body, align=align_t, border=thin,
    )
    section_bar(ws, 10, 8, "2.2 物理／クラウド構成要素")
    table_header(ws, 11, ["構成ID", "リソース", "サービス", "用途", "冗長", "環境", "概算規模", "備考"])
    comps = [
        ["CMP-01", "Web配信", "CloudFront+S3", "SPA・公開アセット", "エッジ", "全", "—", "OAC"],
        ["CMP-02", "APIクラスタ", "ECS Fargate", "REST API", "Multi-AZ", "全", "2–8タスク", "オートスケール"],
        ["CMP-03", "ワーカー", "ECS Fargate", "通知・集計", "Multi-AZ", "全", "1–4", "キュー駆動"],
        ["CMP-04", "DB", "Aurora PostgreSQL", "トランザクションデータ", "Multi-AZ", "全", "r6g.large〜", "Reader付"],
        ["CMP-05", "Redis", "ElastiCache", "キャッシュ・分散ロック", "Multi-AZ", "全", "cache.t4g.m", ""],
        ["CMP-06", "教材桶", "S3", "コンテンツ包", "99.999999999%", "全", "初期100GB", "版プレフィックス"],
        ["CMP-07", "私有サブネット", "VPC", "API/DB配置", "2AZ+", "全", "—", "東京"],
        ["CMP-08", "WAF", "AWS WAF", "OWASP基本ルール", "—", "商用", "—", "CloudFront前面"],
        ["CMP-09", "秘密情報", "Secrets Manager", "DB/API鍵", "—", "全", "—", "ローテーション"],
        ["CMP-10", "DNS/TLS", "Route53+ACM", "名前解決・証明書", "—", "全", "—", ""],
    ]
    for i, t in enumerate(comps):
        table_row(ws, 12 + i, t, alt=i % 2 == 1)

    section_bar(ws, 24, 8, "2.3 環境構成")
    table_header(ws, 25, ["環境", "目的", "データ", "規模比", "個人情報", "接続先人事", "URL例", "備考"])
    envs = [
        ["dev", "開発", "マスキング／ダミー", "0.25", "不可", "モック", "dev.jlms.example", ""],
        ["stg", "結合・受入", "匿名化コピー", "0.5", "原則不可", "検証IF", "stg.jlms.example", "UAT"],
        ["prod", "本番", "本番", "1.0", "可", "本番IF", "jlms.example", ""],
    ]
    for i, t in enumerate(envs):
        table_row(ws, 26 + i, t, alt=i % 2 == 1)

    # ── 機能構成 ──
    ws = wb.create_sheet("04_機能構成")
    set_col_widths(ws, [12, 16, 20, 28, 14, 14, 12, 12])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "3. 機能構成・モジュール")
    table_header(ws, 3, ["モジュール", "サブモジュール", "責務", "主要APIプレフィックス", "要件ID", "配置", "公開", "備考"])
    mods = [
        ["Identity", "Auth", "OIDCコールバック・セッション", "/api/v1/auth", "FR-AUTH-*", "API", "外部", ""],
        ["Identity", "RBAC", "ロール・権限判定", "/api/v1/roles", "FR-AUTH-02", "API", "内部", ""],
        ["Directory", "Users", "ユーザCRUD・招待", "/api/v1/users", "FR-USR-01", "API", "外部", ""],
        ["Directory", "Org", "組織階層", "/api/v1/orgs", "FR-USR-01", "API", "外部", ""],
        ["Directory", "HrSync", "人事差分適用", "（バッチ）", "FR-USR-02", "Worker", "内部", "IF-01"],
        ["Content", "Packs", "教材メタ・版", "/api/v1/contents", "FR-CNT-*", "API", "外部", ""],
        ["Content", "Assets", "署名URL発行", "/api/v1/assets", "FR-CNT-01", "API", "外部", "S3"],
        ["Learning", "Paths", "パス割当・進捗", "/api/v1/paths", "FR-LRN-02", "API", "外部", ""],
        ["Learning", "Events", "学習イベント受信", "/api/v1/events", "FR-LRN-01", "API", "外部", "高頻度"],
        ["Learning", "Assignments", "課題提出", "/api/v1/assignments", "FR-LRN-03", "API", "外部", ""],
        ["Learning", "SRS", "語彙デッキ", "/api/v1/decks", "FR-LRN-04", "API", "外部", "P2"],
        ["Assessment", "Exams", "試験実施", "/api/v1/exams", "FR-ASM-01", "API", "外部", ""],
        ["Assessment", "Grading", "採点WF", "/api/v1/grading", "FR-ASM-02", "API", "外部", ""],
        ["Analytics", "Dashboards", "集計読取モデル", "/api/v1/dashboards", "FR-RPT-*", "API", "外部", ""],
        ["Notify", "Mailer", "通知送信", "（ワーカー）", "FR-NTF-01", "Worker", "内部", "SES"],
        ["Platform", "Audit", "監査ログ", "/api/v1/audit", "FR-AUD-01", "API", "外部", ""],
        ["Platform", "Tenant", "テナント設定", "/api/v1/tenant", "FR-SYS-01", "API", "外部", ""],
    ]
    for i, t in enumerate(mods):
        table_row(ws, 4 + i, t, alt=i % 2 == 1)

    # ── 画面遷移 ──
    ws = wb.create_sheet("05_画面遷移")
    set_col_widths(ws, [14, 16, 16, 18, 20, 14, 14, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "4. 画面遷移設計")
    section_bar(ws, 3, 8, "4.1 学習者サイト遷移（主要）")
    table_header(ws, 4, ["From", "イベント", "To", "条件", "渡す情報", "戻る", "権限", "備考"])
    flows = [
        ["SCR-LOGIN", "SSO成功", "SCR-HOME-L", "ロール=学習者", "session", "—", "認証済", ""],
        ["SCR-HOME-L", "パス選択", "SCR-PATH", "割当あり", "pathId", "Home", "学習者", ""],
        ["SCR-PATH", "教材開始", "SCR-CONTENT", "公開中", "contentId,pos", "Path", "学習者", ""],
        ["SCR-CONTENT", "課題へ", "SCR-ASSIGN", "課題あり", "assignmentId", "Content", "学習者", ""],
        ["SCR-PATH", "試験開始", "SCR-EXAM", "期間内", "examId", "不可（ポリシー）", "学習者", "離脱ルール"],
        ["SCR-EXAM", "提出完了", "SCR-RESULT", "提出成功", "attemptId", "Home", "学習者", "即時/待機"],
        ["SCR-HOME-L", "通知クリック", "対象画面", "権限OK", "deepLink", "—", "学習者", ""],
        ["いずれか", "401/セッション切れ", "SCR-LOGIN", "トークン無効", "returnUrl", "—", "—", ""],
        ["いずれか", "403", "SCR-ERR", "権限不足", "code", "Home", "—", ""],
    ]
    for i, t in enumerate(flows):
        table_row(ws, 5 + i, t, alt=i % 2 == 1, heights=20)

    section_bar(ws, 16, 8, "4.2 講師／管理者サイト遷移（主要）")
    table_header(ws, 17, ["From", "イベント", "To", "条件", "渡す情報", "戻る", "権限", "備考"])
    flows2 = [
        ["SCR-LOGIN", "SSO成功", "SCR-HOME-I", "ロール=講師", "session", "—", "講師", ""],
        ["SCR-HOME-I", "要採点", "SCR-GRADE", "未採点>0", "filter=pending", "Home", "講師", ""],
        ["SCR-HOME-I", "教材管理", "SCR-CNT-ADM", "権限", "—", "Home", "講師/管理者", ""],
        ["SCR-CNT-ADM", "公開実行", "SCR-CNT-ADM", "バリデーションOK", "version", "—", "公開権限", "確認モーダル"],
        ["SCR-HOME-I", "クラス進捗", "SCR-CLASS", "担当クラス", "classId", "Home", "講師", ""],
        ["SCR-USR", "招待送信", "SCR-USR", "メール妥当", "inviteId", "—", "管理者", ""],
        ["SCR-AUDIT", "エクスポート", "ダウンロード", "期間指定", "jobId", "—", "SysAdmin", "非同期"],
    ]
    for i, t in enumerate(flows2):
        table_row(ws, 18 + i, t, alt=i % 2 == 1, heights=20)

    # ── 画面設計 ──
    ws = wb.create_sheet("06_画面設計")
    set_col_widths(ws, [12, 16, 14, 28, 18, 14, 12, 12])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "5. 画面設計（レイアウト方針・画面仕様概要）")
    section_bar(ws, 3, 8, "5.1 UI方針")
    merge_write(
        ws, "A4", "H6",
        "・デスクトップ優先＋モバイルで主要学習フロー完結（Responsive）。\n"
        "・情報設計は「次にやること」を最上位に置く（学習者ホーム）。\n"
        "・共通レイアウト：ヘッダー（テナントロゴ／言語／通知／ユーザ）＋サイドナビ（ロール別）＋コンテンツ。\n"
        "・アクセシビリティ：コントラスト、キーボード操作、フォーカス可視化、代替テキスト（AA目標）。\n"
        "・文言：UIは日英切替。エラーメッセージはコード＋人間可読文（詳細設計のメッセージ一覧に従う）。",
        fill=fill_white, font=font_body, align=align_t, border=thin,
    )
    section_bar(ws, 8, 8, "5.2 画面仕様概要（代表）")
    table_header(ws, 9, ["画面ID", "画面名", "レイアウト", "主要コンポーネント", "API", "イベント", "権限", "NFR"])
    screens = [
        ["SCR-HOME-L", "学習者ホーム", "2カラム", "次の教材カード、進捗リング、通知", "GET /dashboards/learner", "表示", "学習者", "p95≦2s"],
        ["SCR-CONTENT", "教材ビューア", "コンテンツ全幅", "ビューア、進捗バー、しおり", "POST /events", "heartbeat", "学習者", "安定通信"],
        ["SCR-EXAM", "試験実施", "集中レイアウト", "問題、タイマ、ナビ、提出", "POST /exams/{id}/answers", "自動保存", "学習者", "p95≦1s"],
        ["SCR-GRADE", "採点WB", "リスト+詳細", "提出一覧、採点フォーム", "GET/POST /grading", "保存/差戻し", "講師", ""],
        ["SCR-CLASS", "クラス進捗", "テーブル+フィルタ", "遅延ハイライト、CSV", "GET /dashboards/class", "エクスポート", "講師", ""],
        ["SCR-CNT-ADM", "教材管理", "テーブル+ドロワ", "版、公開予約、プレビュー", "CRUD /contents", "公開", "講師+", ""],
        ["SCR-USR", "ユーザ管理", "テーブル+モーダル", "招待、ロール、無効化", "CRUD /users", "招待", "管理者", "監査"],
        ["SCR-AUDIT", "監査ログ", "検索フォーム+表", "期間、アクター、エクスポート", "GET /audit", "検索", "SysAdmin", ""],
    ]
    for i, t in enumerate(screens):
        table_row(ws, 10 + i, t, alt=i % 2 == 1, heights=28)

    section_bar(ws, 20, 8, "5.3 項目定義例：SCR-EXAM（抜粋）")
    table_header(ws, 21, ["項目ID", "項目名", "種別", "必須", "初期値", "バリデーション", "権限", "備考"])
    items = [
        ["E-01", "残り時間", "表示", "—", "試験設定", "—", "学習者", "1秒更新"],
        ["E-02", "問題文", "表示", "—", "API", "—", "学習者", "HTMLサニタイズ"],
        ["E-03", "選択肢", "ラジオ", "○", "未選択", "1つ選択", "学習者", "四択時"],
        ["E-04", "記述回答", "テキスト", "○", "空", "1–2000文字", "学習者", "記述時"],
        ["E-05", "音声再生", "コントロール", "—", "停止", "再生回数制限可", "学習者", "聴解"],
        ["E-06", "一時保存", "ボタン", "—", "—", "—", "学習者", "自動でも実行"],
        ["E-07", "提出", "ボタン", "—", "無効→有効", "未回答確認ダイアログ", "学習者", "二重送信防止"],
        ["E-08", "設問ナビ", "リスト", "—", "現在位置", "—", "学習者", "フラグ機能"],
    ]
    for i, t in enumerate(items):
        table_row(ws, 22 + i, t, alt=i % 2 == 1)

    # ── 帳票 ──
    ws = wb.create_sheet("07_帳票設計")
    set_col_widths(ws, [10, 18, 12, 28, 14, 14, 12, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "6. 帳票設計")
    table_header(ws, 3, ["帳票ID", "名称", "形式", "レイアウト概要", "生成方式", "非同期", "暗号化", "配信"])
    reports = [
        ["RPT-01", "個人学習履歴", "PDF", "表紙+期間+履歴表+技能レーダー", "サーバ生成", "小容量同期", "HTTPS", "DL"],
        ["RPT-02", "クラス進捗一覧", "XLSX", "1行1学習者、進捗%/最終学習日/遅延", "ストリーム", "可", "HTTPS", "DL"],
        ["RPT-03", "組織月次", "PDF", "サマリKPI+部署別表+前月比", "バッチ/手動", "必須", "HTTPS", "DL/メール"],
        ["RPT-04", "試験結果明細", "CSV", "attempt単位の明細", "クエリ出力", "可", "HTTPS", "DL"],
        ["RPT-05", "監査ログ", "CSV", "時系列ログ", "非同期ジョブ", "必須", "ZIP+AES", "DL"],
        ["RPT-06", "同期エラー", "CSV", "失敗行+理由コード", "バッチ成果物", "—", "HTTPS", "管理画面"],
    ]
    for i, t in enumerate(reports):
        table_row(ws, 4 + i, t, alt=i % 2 == 1, heights=24)

    section_bar(ws, 12, 8, "6.1 RPT-03 組織月次　項目（抜粋）")
    table_header(ws, 13, ["項番", "項目", "型", "算出ロジック", "単位", "個人情報", "表示条件", "備考"])
    ritems = [
        ["1", "対象年月", "YYYY-MM", "パラメータ", "—", "否", "常時", ""],
        ["2", "アクティブ学習者数", "number", "期間内イベントありのユニークユーザ", "人", "否", "常時", ""],
        ["3", "平均進捗率", "number", "Enrollment.progress平均", "%", "否", "常時", "小数1桁"],
        ["4", "遅延者数", "number", "期限超過かつ進捗<100%", "人", "否", "常時", ""],
        ["5", "部署別進捗", "表", "OrgUnit単位集計", "%", "否", "常時", ""],
        ["6", "試験平均点", "number", "期間内Attempt.score平均", "点", "否", "試験あり時", ""],
    ]
    for i, t in enumerate(ritems):
        table_row(ws, 14 + i, t, alt=i % 2 == 1)

    # ── データモデル ──
    ws = wb.create_sheet("08_データモデル")
    set_col_widths(ws, [16, 16, 14, 14, 14, 14, 14, 18])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "7. 論理データモデル")
    section_bar(ws, 3, 8, "7.1 エンティティ関連（文章ER）")
    merge_write(
        ws, "A4", "H8",
        "Tenant 1─* User\n"
        "Tenant 1─* OrgUnit（階層：OrgUnit 1─* OrgUnit）\n"
        "User *─* OrgUnit（所属）\n"
        "Tenant 1─* ContentPack 1─* ContentVersion\n"
        "Tenant 1─* LearningPath 1─* LearningPathItem（ContentVersion / Assessment を参照）\n"
        "User *─* LearningPath（Enrollment）1─* LearningEvent\n"
        "Assessment 1─* Question / Attempt 1─* Answer\n"
        "User 1─* Attempt / User 1─* Notification / User 1─* AuditLog（actor）",
        fill=fill_light, font=font_body, align=align_t, border=thin,
    )
    section_bar(ws, 10, 8, "7.2 キー設計方針")
    table_header(ws, 11, ["方針ID", "内容", "適用", "理由", "例外", "詳細設計", "状態", "備考"])
    keys = [
        ["KEY-01", "サロゲートPKはUUID v7", "全テーブル", "分散生成・時系列", "マスタコード表", "DD参照", "確定", ""],
        ["KEY-02", "全業務表にtenant_id", "テナントデータ", "分離・強制フィルタ", "プラットフォーム表", "RLS検討", "確定", ""],
        ["KEY-03", "ビジネスキーユニーク", "email+tenant等", "同期・冪等", "—", "ユニーク制約", "確定", ""],
        ["KEY-04", "論理削除（deleted_at）", "マスタ系", "監査・復元", "イベント・ログ", "—", "確定", ""],
        ["KEY-05", "楽観ロック（version）", "更新競合あり表", "同時更新", "追記のみ表", "—", "確定", ""],
    ]
    for i, t in enumerate(keys):
        table_row(ws, 12 + i, t, alt=i % 2 == 1, heights=20)

    # ── テーブル一覧 ──
    ws = wb.create_sheet("09_テーブル一覧")
    set_col_widths(ws, [6, 22, 28, 12, 12, 12, 14, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "8. テーブル一覧（論理）")
    table_header(ws, 3, ["#", "テーブル名", "説明", "個人情報", "想定件数", "保持", "モジュール", "要件"])
    tables = [
        ["1", "tenants", "テナント", "否", "小", "契約+1年", "Platform", "FR-SYS-01"],
        ["2", "users", "ユーザ", "含", "2,000+", "退職+5年", "Directory", "FR-USR-*"],
        ["3", "org_units", "組織", "否", "200", "存続", "Directory", "FR-USR-01"],
        ["4", "user_org_memberships", "所属", "含", "4,000", "退職+5年", "Directory", "FR-USR-01"],
        ["5", "roles", "ロール定義", "否", "小", "存続", "Identity", "FR-AUTH-02"],
        ["6", "user_roles", "ロール付与", "含", "2,000", "退職+5年", "Identity", "FR-AUTH-02"],
        ["7", "content_packs", "教材包", "否", "500", "7年", "Content", "FR-CNT-01"],
        ["8", "content_versions", "教材版", "否", "2,000", "7年", "Content", "FR-CNT-02"],
        ["9", "learning_paths", "学習パス", "否", "50", "存続", "Learning", "FR-LRN-02"],
        ["10", "learning_path_items", "パス構成要素", "否", "1,000", "存続", "Learning", "FR-LRN-02"],
        ["11", "enrollments", "受講", "含", "4,000", "5年", "Learning", "FR-LRN-02"],
        ["12", "learning_events", "学習イベント", "含", "百万", "3年", "Learning", "FR-LRN-01"],
        ["13", "assignments", "課題定義", "否", "500", "存続", "Learning", "FR-LRN-03"],
        ["14", "submissions", "提出", "含", "2万", "5年", "Learning", "FR-LRN-03"],
        ["15", "assessments", "試験定義", "否", "100", "存続", "Assessment", "FR-ASM-01"],
        ["16", "questions", "設問", "否", "5,000", "存続", "Assessment", "FR-ASM-01"],
        ["17", "attempts", "受験", "含", "2万", "5年", "Assessment", "FR-ASM-*"],
        ["18", "answers", "回答", "含", "20万", "5年", "Assessment", "FR-ASM-*"],
        ["19", "notifications", "通知", "含", "50万", "1年", "Notify", "FR-NTF-01"],
        ["20", "audit_logs", "監査", "含", "100万", "3年", "Platform", "FR-AUD-01"],
        ["21", "hr_sync_runs", "同期実行", "含", "数千", "1年", "Directory", "FR-IF-01"],
        ["22", "hr_sync_errors", "同期エラー行", "含", "数万", "90日", "Directory", "FR-IF-01"],
        ["23", "feature_flags", "機能フラグ", "否", "小", "存続", "Platform", "AP-06"],
        ["24", "message_catalogs", "メッセージ", "否", "小", "存続", "Platform", "UI"],
    ]
    for i, t in enumerate(tables):
        table_row(ws, 4 + i, t, alt=i % 2 == 1)

    # ── IF設計 ──
    ws = wb.create_sheet("10_IF設計")
    set_col_widths(ws, [10, 16, 12, 14, 28, 12, 12, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "9. 外部インターフェース設計")
    section_bar(ws, 3, 8, "9.1 IF一覧（基本設計レベル）")
    table_header(ws, 4, ["IF-ID", "名称", "プロトコル", "認証", "エンドポイント概要", "タイムアウト", "再試行", "冪等"])
    ifs = [
        ["IF-01", "人事差分", "HTTPS REST", "mTLS+API Key", "GET /hr/v1/users?updatedSince=", "30s", "3回", "○（updatedAt）"],
        ["IF-02", "OIDC", "HTTPS", "OAuth2/OIDC", "Authorize/Token/JWKS", "10s", "—", "—"],
        ["IF-03", "SES", "AWS SDK", "IAM Role", "SendEmail", "10s", "3回", "○ MessageId"],
        ["IF-04", "S3", "AWS SDK", "IAM Role", "Put/Get/Presign", "60s", "SDK標準", "○ key"],
        ["IF-05", "CloudFront", "HTTPS", "署名URL", "教材配信", "—", "—", "—"],
    ]
    for i, t in enumerate(ifs):
        table_row(ws, 5 + i, t, alt=i % 2 == 1, heights=22)

    section_bar(ws, 12, 8, "9.2 内部API共通仕様")
    table_header(ws, 13, ["項目", "仕様", "例", "必須", "備考", "", "", ""])
    api_common = [
        ["ベースURL", "/api/v1", "https://jlms.example/api/v1", "○", "バージョニング", "", "", ""],
        ["認証", "Bearer JWT", "Authorization: Bearer …", "○", "除：ヘルス", "", "", ""],
        ["テナント", "JWT claim + 必要時ヘッダ", "tid", "○", "不一致は403", "", "", ""],
        ["相関ID", "X-Request-Id", "UUID", "○", "ログ横断", "", "", ""],
        ["日時", "ISO-8601 UTC", "2026-08-01T15:00:00Z", "○", "UIはローカル表示", "", "", ""],
        ["エラー形式", "application/problem+json", "type,title,status,detail,code", "○", "メッセージコード連携", "", "", ""],
        ["ページング", "cursor or offset", "limit≦100", "一覧", "既定 limit=20", "", "", ""],
        ["楽観ロック", "If-Match / version", "version=3", "更新系", "409 Conflict", "", "", ""],
    ]
    for i, t in enumerate(api_common):
        table_row(ws, 14 + i, t, alt=i % 2 == 1)

    section_bar(ws, 24, 8, "9.3 代表API一覧（抜粋）")
    table_header(ws, 25, ["API-ID", "Method", "Path", "概要", "画面/IF", "認可", "SLIp95", "備考"])
    apis = [
        ["API-001", "GET", "/paths/me", "自分の学習パス一覧", "SCR-HOME-L", "Learner", "2s", ""],
        ["API-002", "POST", "/events", "学習イベント一括送信", "SCR-CONTENT", "Learner", "1s", "最大50件"],
        ["API-003", "POST", "/exams/{id}/start", "試験開始", "SCR-EXAM", "Learner", "1s", ""],
        ["API-004", "PUT", "/exams/attempts/{id}", "回答保存", "SCR-EXAM", "Learner", "1s", "冪等"],
        ["API-005", "POST", "/exams/attempts/{id}/submit", "提出", "SCR-EXAM", "Learner", "1s", ""],
        ["API-006", "GET", "/grading/inbox", "採点inbox", "SCR-GRADE", "Instructor", "2s", ""],
        ["API-007", "POST", "/grading/{id}/score", "採点確定", "SCR-GRADE", "Instructor", "1s", "監査"],
        ["API-008", "POST", "/contents/{id}/publish", "教材公開", "SCR-CNT-ADM", "Instructor+", "2s", ""],
        ["API-009", "POST", "/users/invitations", "招待", "SCR-USR", "Admin", "2s", ""],
        ["API-010", "GET", "/dashboards/class/{id}", "クラス進捗", "SCR-CLASS", "Instructor", "2s", ""],
        ["API-011", "GET", "/audit", "監査検索", "SCR-AUDIT", "SysAdmin", "3s", ""],
        ["API-012", "GET", "/health", "ヘルス", "監視", "公開", "0.2s", "認証不要"],
    ]
    for i, t in enumerate(apis):
        table_row(ws, 26 + i, t, alt=i % 2 == 1)

    # ── バッチ ──
    ws = wb.create_sheet("11_バッチ設計")
    set_col_widths(ws, [12, 18, 14, 14, 20, 14, 12, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "10. バッチ設計")
    table_header(ws, 3, ["バッチID", "名称", "スケジュール", "実行時間目標", "入力", "出力", "失敗時", "監視"])
    batches = [
        ["BAT-01", "人事差分同期", "毎日 02:00 JST", "≦30分", "IF-01", "users更新+エラーCSV", "リトライ後アラート", "必須"],
        ["BAT-02", "月次組織レポート", "毎月1日 03:00", "≦20分", "集計テーブル", "PDF/S3+メール", "再実行可", "必須"],
        ["BAT-03", "リマインド抽出", "毎日 09:00 JST", "≦15分", "enrollment/期限", "通知キュー", "部分再実行", "必須"],
        ["BAT-04", "学習イベント集計", "毎時", "≦10分", "events", "ダッシュボード用集計", "次時間で追いつき", "必須"],
        ["BAT-05", "監査ログアーカイブ", "毎日 04:00", "≦30分", "audit_logs", "S3 Glacier", "翌日再実行", ""],
        ["BAT-06", "期限切れ署名掃除", "毎日 01:00", "≦10分", "一時表", "削除", "無視可", ""],
    ]
    for i, t in enumerate(batches):
        table_row(ws, 4 + i, t, alt=i % 2 == 1, heights=22)

    section_bar(ws, 12, 8, "10.1 BAT-01 処理フロー")
    merge_write(
        ws, "A13", "H16",
        "1) 前回成功カーソル（updatedAt）取得 → 2) IF-01ページング取得 → 3) 行バリデーション\n"
        "4) 件数閾値チェック（例: 変更率20%超で停止）→ 5) UPSERT/無効化（トランザクション単位はチャンク）\n"
        "6) hr_sync_runs にサマリ記録 → 7) エラー行を hr_sync_errors へ → 8) 運用通知（成功/失敗）\n"
        "再実行：同一カーソルで冪等。部分成功時は失敗行のみ是正後に再同期。",
        fill=fill_white, font=font_body, align=align_t, border=thin,
    )

    # ── セキュリティ ──
    ws = wb.create_sheet("12_セキュリティ")
    set_col_widths(ws, [12, 16, 28, 16, 14, 12, 12, 12])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "11. セキュリティ設計")
    table_header(ws, 3, ["統制ID", "領域", "対策", "実装方針", "要件", "検証", "優先", "状態"])
    secs = [
        ["SEC-01", "認証", "OIDC SSO + 管理者MFA", "Entra ID条件付きアクセス連携", "NFR-09", "設定監査", "P1", "確定"],
        ["SEC-02", "認可", "RBAC＋リソース所属チェック", "サーバ側強制（UI隠蔽のみに依存しない）", "FR-AUTH-02", "権限試験", "P1", "確定"],
        ["SEC-03", "テナント分離", "tenant_id強制・試験で越権検知", "リポジトリ層で自動付与", "FR-SYS-01", "セキュリティ試験", "P1", "確定"],
        ["SEC-04", "暗号", "TLS1.2+ / DB・S3暗号化", "AWS KMS CMK", "NFR-06/07", "構成スキャン", "P1", "確定"],
        ["SEC-05", "入力", "バリデーション・サニタイズ", "クラスバリデーション+HTML Purifier", "OWASP", "SAST/DAST", "P1", "確定"],
        ["SEC-06", "ファイル", "教材タイプ制限・ウイルススキャン", "拡張子+MIME+ClamAV相当", "セキュポリ", "結合", "P1", "確定"],
        ["SEC-07", "監査", "重要操作の監査ログ", "改ざん防止（追記専用）", "FR-AUD-01", "レビュー", "P1", "確定"],
        ["SEC-08", "秘密情報", "Secrets Manager、コード埋め込み禁止", "CIシークレットスキャン", "セキュポリ", "監査", "P1", "確定"],
        ["SEC-09", "セッション", "短命アクセストークン+刷新", "Redisブラックリスト（ログアウト）", "NFR", "試験", "P1", "確定"],
        ["SEC-10", "個人情報", "国外転送禁止・最小化", "東京リージョン固定、項目最小化", "NFR-19", "設計レビュー", "P1", "確定"],
    ]
    for i, t in enumerate(secs):
        table_row(ws, 4 + i, t, alt=i % 2 == 1, heights=24)

    # ── 運用監視 ──
    ws = wb.create_sheet("13_運用監視")
    set_col_widths(ws, [12, 18, 28, 14, 14, 14, 12, 12])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "12. 運用・監視設計")
    table_header(ws, 3, ["監視ID", "対象", "条件", "重大度", "通知先", "ランブック", "SLO", "備考"])
    mons = [
        ["MON-01", "API 5xx率", "5分で>1%", "SEV1", "Pager", "RB-API-5XX", "99.9%", ""],
        ["MON-02", "p95レイテンシ", "主要エンドポイント>2s", "SEV2", "Slack", "RB-LAT", "NFR-01", ""],
        ["MON-03", "DB CPU/接続", "CPU>80% 10分", "SEV2", "Pager", "RB-DB", "—", ""],
        ["MON-04", "BAT-01失敗", "ジョブFailed", "SEV1", "Pager", "RB-HR-SYNC", "NFR-18", ""],
        ["MON-05", "証明書期限", "残14日", "SEV3", "メール", "RB-CERT", "—", ""],
        ["MON-06", "ディスク/S3エラー", "Put失敗率上昇", "SEV2", "Slack", "RB-S3", "—", ""],
        ["MON-07", "外形監視", "ログイン+ホーム取得失敗", "SEV1", "Pager", "RB-SYNTH", "NFR-04", "1分間隔"],
    ]
    for i, t in enumerate(mons):
        table_row(ws, 4 + i, t, alt=i % 2 == 1)

    section_bar(ws, 13, 8, "12.1 リリース・変更管理（概要）")
    merge_write(
        ws, "A14", "H16",
        "・デプロイ：Blue/Green（ECS）。DBマイグレーションは前方互換を原則とし、拡大→デプロイ→縮小の順。\n"
        "・Feature Flagで教材AI採点等の未完成機能を隔離。\n"
        "・本番変更はRFC＋承認。緊急は事後RFC（24時間以内）。",
        fill=fill_white, font=font_body, align=align_t, border=thin,
    )

    # ── トレーサビリティ ──
    ws = wb.create_sheet("14_トレーサビリティ")
    set_col_widths(ws, [14, 14, 16, 16, 14, 14, 12, 14])
    ws.sheet_view.showGridLines = False
    header_bar(ws, 1, 8, "13. 要件→基本設計トレーサビリティ（抜粋）")
    table_header(ws, 3, ["要件ID", "モジュール", "画面ID", "API/バッチ", "テーブル", "IF", "試験観点", "備考"])
    tm = [
        ["FR-AUTH-01", "Identity", "SCR-LOGIN", "OIDC", "users", "IF-02", "認証", ""],
        ["FR-USR-02", "Directory.HrSync", "SCR-BATCH", "BAT-01", "hr_sync_*", "IF-01", "結合", ""],
        ["FR-CNT-02", "Content", "SCR-CNT-ADM", "API-008", "content_versions", "IF-04", "機能", ""],
        ["FR-LRN-01", "Learning.Events", "SCR-CONTENT", "API-002", "learning_events", "—", "機能", ""],
        ["FR-LRN-02", "Learning.Paths", "SCR-PATH", "API-001", "enrollments", "—", "機能", ""],
        ["FR-ASM-01", "Assessment", "SCR-EXAM", "API-003〜005", "attempts/answers", "—", "機能", ""],
        ["FR-ASM-02", "Assessment.Grading", "SCR-GRADE", "API-006/007", "submissions", "—", "機能", ""],
        ["FR-RPT-01", "Analytics", "SCR-CLASS", "API-010", "集計", "RPT-02", "機能", ""],
        ["FR-NTF-01", "Notify", "—", "BAT-03", "notifications", "IF-03", "結合", ""],
        ["FR-AUD-01", "Platform.Audit", "SCR-AUDIT", "API-011", "audit_logs", "RPT-05", "セキュリティ", ""],
        ["NFR-01", "APIクラスタ", "主要画面", "—", "—", "—", "性能", ""],
        ["NFR-04", "全体", "—", "MON-07", "—", "—", "信頼性", ""],
        ["NFR-06", "データ層", "—", "—", "全表", "—", "セキュリティ", ""],
    ]
    for i, t in enumerate(tm):
        table_row(ws, 4 + i, t, alt=i % 2 == 1)

    path = OUT_DIR / "02_基本設計書_JLMS.xlsx"
    wb.save(path)
    return path
